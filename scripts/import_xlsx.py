"""One-time migration: read the Google Sheets export and write data/seed.json.

Usage:  python -I scripts/import_xlsx.py "My Cashflow Tracker.xlsx" data/seed.json

All money is written as integer cents and all dates as YYYY-MM-DD strings.
The output contains real financial data - it lives under data/, which is gitignored.
"""
import calendar
import datetime as dt
import json
import os
import sys

import openpyxl

BILL_ROWS = range(18, 35)      # Setup!A18:A34
INCOME_ROWS = range(37, 42)    # Setup!A37:A41
ONEOFF_ROWS = range(45, 60)    # Setup!A45:A59
GRID_COLS = [("M", "N"), ("O", "P"), ("Q", "R")]  # (amount, paid) for M1, M2, M3

# Codes that exist only to tag transactions; they never generate scheduled events.
UNSCHEDULED = {"XFER", "MISC", "BYE"}


def cents(v):
    return int(round(float(v) * 100))


def iso(v):
    if isinstance(v, dt.datetime):
        v = v.date()
    return v.isoformat()


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return dt.date(d.year + y, m + 1, 1)


def due_date(month_start, day):
    last = calendar.monthrange(month_start.year, month_start.month)[1]
    return month_start.replace(day=min(int(day), last))


def main(src, dst):
    global EXPORT_DATE
    # Cached formula values (e.g. TODAY()) were computed when the file was exported.
    EXPORT_DATE = dt.date.fromtimestamp(os.path.getmtime(src)).isoformat()
    wb = openpyxl.load_workbook(src, data_only=True)
    setup, hist = wb["Setup"], wb["History"]
    v = lambda ref: setup[ref].value

    window_start = v("C4").date()
    pay_anchor = v("C9").date()

    settings = {
        "bufferFloorCents": cents(v("C8")),
        "savingsFloorCents": cents(v("C87")),
        "horizonDays": int(v("C6")),
        "lookbackDays": int(v("C12")),
        "estimateMode": "worst" if v("C13") == "Worst" else "average",
        "topUpRoundingCents": 5000,
        "savingsDrawCodes": ["XFER"],
    }

    bills, sort = [], 0
    for rows, direction in ((BILL_ROWS, "out"), (INCOME_ROWS, "in")):
        for r in rows:
            code = setup[f"A{r}"].value
            if not code:
                continue
            code = str(code).strip()
            sort += 1
            mode = setup[f"G{r}"].value or "Manual"
            manual = setup[f"H{r}"].value
            floor = setup[f"U{r}"].value
            bill = {
                "code": code,
                "name": setup[f"B{r}"].value,
                "category": setup[f"C{r}"].value or "Uncategorized",
                "direction": direction,
                "active": setup[f"F{r}"].value == "Y",
                "schedule": {"kind": "monthly", "day": int(setup[f"D{r}"].value or 1)},
                "amountMode": "auto" if mode == "Auto" else "fixed",
                "amountCents": cents(manual) if isinstance(manual, (int, float)) else 0,
                "minCents": cents(floor) if isinstance(floor, (int, float)) and floor else None,
                "notes": setup[f"S{r}"].value or "",
                "sort": sort,
            }
            if code in UNSCHEDULED:
                bill["schedule"] = {"kind": "none"}
                bill["active"] = False
            if code == "PAY":
                # The sheet generated paychecks from the anchor, not from this row.
                bill.update(
                    active=True,
                    schedule={"kind": "interval", "anchor": iso(pay_anchor), "everyDays": int(v("C11"))},
                    amountMode="fixed",
                    amountCents=cents(v("C10")),
                    notes="Bi-weekly paycheck (imported from Setup payday anchor).",
                )
            bills.append(bill)

    by_code = {b["code"]: b for b in bills}

    transactions, overrides = [], []
    for r in hist.iter_rows(min_row=3, values_only=True):
        if not r[0]:
            continue
        transactions.append({
            "date": iso(r[0]), "code": r[2], "description": r[1], "direction": r[3],
            "amountCents": cents(r[4] or 0), "source": "sheet-history", "note": "",
        })

    # Setup M1/M2/M3 grid: paid -> transaction on the due date; unpaid amount -> known override.
    for r in list(BILL_ROWS) + list(INCOME_ROWS):
        code = setup[f"A{r}"].value
        if not code or code not in by_code:
            continue
        bill = by_code[code]
        if not bill["active"] or bill["schedule"]["kind"] != "monthly":
            continue
        for k, (amt_col, pd_col) in enumerate(GRID_COLS):
            amt, paid = setup[f"{amt_col}{r}"].value, setup[f"{pd_col}{r}"].value
            if not isinstance(amt, (int, float)):
                continue
            month = add_months(window_start, k)
            due = due_date(month, bill["schedule"]["day"])
            # A future-dated "paid" tick means the amount is certain, not that money moved yet.
            if (paid is True or paid == "Y") and iso(due) <= EXPORT_DATE:
                transactions.append({
                    "date": iso(due), "code": code,
                    "description": bill["name"], "direction": bill["direction"],
                    "amountCents": cents(amt), "source": "sheet-confirmed", "note": "",
                })
            else:
                overrides.append({"code": code, "period": month.strftime("%Y-%m"), "amountCents": cents(amt)})

    planned = []
    for r in ONEOFF_ROWS:
        code, date = setup[f"A{r}"].value, setup[f"B{r}"].value
        if not code or not date:
            continue
        amount, actual, paid = setup[f"C{r}"].value, setup[f"G{r}"].value, setup[f"H{r}"].value
        amt = actual if isinstance(actual, (int, float)) else amount
        bill = by_code.get(code)
        planned.append({
            "date": iso(date), "code": code,
            "description": bill["name"] if bill else code,
            "direction": bill["direction"] if bill else "out",
            "amountCents": cents(amt or 0), "note": setup[f"I{r}"].value or "",
            "status": "done" if paid is True or paid == "Y" else "planned",
        })

    snapshots = [
        {"account": "checking", "date": iso(window_start), "amountCents": cents(v("C7"))},
        # The sheet does not record when savings was last reconciled; assume the export date.
        {"account": "savings", "date": EXPORT_DATE, "amountCents": cents(v("C86"))},
    ]

    out = {"settings": settings, "bills": bills, "transactions": transactions,
           "overrides": overrides, "planned": planned, "snapshots": snapshots}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1)

    # The sheet's own Ledger results, used by tests/parity.test.ts.
    ledger = wb["Ledger"]
    end_of_day = {}
    for row in ledger.iter_rows(min_row=12, max_row=106, max_col=8, values_only=True):
        if isinstance(row[0], dt.datetime) and isinstance(row[7], (int, float)):
            end_of_day[iso(row[0])] = cents(row[7])
    expected = {
        "windowStart": iso(window_start),
        "windowEnd": iso(window_start + dt.timedelta(days=int(v("C6")))),
        "balanceToday": {"date": EXPORT_DATE, "balanceCents": cents(ledger["B3"].value)},
        "endOfDay": end_of_day,
    }
    with open(dst.replace("seed.json", "sheet-ledger.json"), "w", encoding="utf-8") as f:
        json.dump(expected, f, indent=1)
    print(f"bills={len(bills)} transactions={len(transactions)} overrides={len(overrides)} "
          f"planned={len(planned)} -> {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
