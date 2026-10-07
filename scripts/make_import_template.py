"""Build the import template that Settings → Import offers for download.

Usage:  python -I scripts/make_import_template.py
Writes src/web/public/cashflow-import-template.xlsx. The columns must match what
src/server/importSheet.ts reads; tests/importSheet.test.ts imports this file.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = 'src/web/public/cashflow-import-template.xlsx'
# Keep in step with CATEGORIES in src/engine/types.ts.
CATEGORIES = ['Housing', 'Support', 'Credit Card', 'Utility', 'Insurance', 'Taxes',
              'Income', 'Transfer', 'Childcare', 'Savings draw', 'Uncategorized']
OFTEN = ['Monthly', 'Every 2 weeks', 'Every week', 'Every 4 weeks', 'One-off type']

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2A78D6')
EXAMPLE = Font(italic=True, color='7A7974')
MONEY = '"$"#,##0.00'
DATE = 'yyyy-mm-dd'


def sheet(wb, title, headers, widths, rows, formats):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for r in rows:
        ws.append(r)
        for c in ws[ws.max_row]:
            c.font = EXAMPLE
    for col, fmt in formats.items():
        for r in range(2, 1001):
            ws[f'{col}{r}'].number_format = fmt
    ws.freeze_panes = 'A2'
    return ws


def pick(ws, col, options):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    dv.add(f'{col}2:{col}1000')
    ws.add_data_validation(dv)


wb = Workbook()
start = wb.active
start.title = 'Start here'
lines = [
    ('Cashflow import template', True),
    ('', False),
    ('Fill in the three tabs, then in the app open Settings → Import from a spreadsheet and choose this file.', False),
    ('You see everything that will be added before anything is saved. Importing the same file twice adds nothing new.', False),
    ('The grey rows are examples: replace or delete them.', False),
    ('', False),
    ('Bills: everything that repeats, in or out (mortgage, paychecks, card payments, utilities).', True),
    ('  Name: what you call it. History rows refer to it by this name.', False),
    ('  Money in or out: Out for bills, In for paychecks and other income.', False),
    ('  Category: pick from the list.', False),
    ('  Amount: the usual amount. Amount varies? Yes means the app estimates it from your History; the amount is used until there is some.', False),
    ('  How often: Monthly (fill in Day of month), Every week / 2 weeks / 4 weeks (fill in Next date: any date it happens),', False),
    ('    or One-off type: never scheduled, just a label for history and one-offs (e.g. "Transfer from savings", "Venmo").', False),
    ('  Code (optional): a short permanent ID like MORT. Left blank, one is made from the name.', False),
    ('', False),
    ('History: past payments and deposits, one row each. Optional, but auto-estimates need it.', True),
    ('  Item: a name from the Bills tab (or already in the app). Anything else becomes a new One-off type;', False),
    ('    for those, fill in Money in or out.', False),
    ('  Amount: positive. The Bills tab says whether it was in or out.', False),
    ('', False),
    ('Balances: at least one checking balance, so the forecast has somewhere to start.', True),
    ('  The forecast runs from the latest checking balance. History dated after it counts against it.', False),
    ('', False),
    ('Tips', True),
    ('  From Google Sheets: File → Download → Microsoft Excel (.xlsx).', False),
    ('  Dates can be real dates or text like 2026-10-01 or 10/1/2026. Amounts can include $ and commas.', False),
    ('  Keep this file private: it holds your financial data.', False),
]
for text, bold in lines:
    start.append([text])
    start.cell(row=start.max_row, column=1).font = Font(bold=bold, size=14 if text == lines[0][0] else 11)
start.column_dimensions['A'].width = 120
start['A3'].alignment = Alignment(wrap_text=False)

bills = sheet(
    wb, 'Bills',
    ['Name', 'Money in or out', 'Category', 'Amount', 'Amount varies?', 'How often', 'Day of month', 'Next date', 'Code'],
    [26, 16, 16, 12, 15, 16, 13, 13, 10],
    [
        ['Mortgage', 'Out', 'Housing', 1850, 'No', 'Monthly', 1, None, 'MORT'],
        ['Paycheck', 'In', 'Income', 2400, 'No', 'Every 2 weeks', None, date(2026, 10, 2), 'PAY'],
        ['Electric', 'Out', 'Utility', 140, 'Yes', 'Monthly', 18, None, None],
        ['Visa card', 'Out', 'Credit Card', 900, 'Yes', 'Monthly', 25, None, None],
        ['Transfer from savings', 'In', 'Savings draw', 0, 'No', 'One-off type', None, None, 'XFER'],
    ],
    {'D': MONEY, 'H': DATE},
)
pick(bills, 'B', ['In', 'Out'])
pick(bills, 'C', CATEGORIES)
pick(bills, 'E', ['Yes', 'No'])
pick(bills, 'F', OFTEN)

hist = sheet(
    wb, 'History',
    ['Date', 'Item', 'Amount', 'Money in or out', 'Note'],
    [13, 26, 12, 16, 36],
    [
        [date(2026, 9, 1), 'Mortgage', 1850, None, None],
        [date(2026, 9, 18), 'Electric', 162.37, None, 'AC ran a lot'],
        [date(2026, 9, 25), 'Visa card', 1043.18, None, None],
        [date(2026, 9, 18), 'Paycheck', 2400, None, None],
        [date(2026, 9, 20), 'Plumber', 275, 'Out', 'Kitchen sink'],
    ],
    {'A': DATE, 'C': MONEY},
)
pick(hist, 'D', ['In', 'Out'])

bal = sheet(
    wb, 'Balances',
    ['Account', 'Date', 'Balance'],
    [12, 13, 14],
    [
        ['Checking', date(2026, 10, 1), 3250],
        ['Savings', date(2026, 10, 1), 15000],
    ],
    {'B': DATE, 'C': MONEY},
)
pick(bal, 'A', ['Checking', 'Savings'])

wb.save(OUT)
print(f'Wrote {OUT}')
